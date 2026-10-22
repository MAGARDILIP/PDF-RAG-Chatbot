"""
Mini Project Synopsis Generator
Generates the mini project synopsis .docx for MCA course MCA33EL03
Title: AI-Powered Document Query System
Student: Dilip Magar | PRN: 125M1H055
Guide: Dr. Anjana Arakerimath
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
import os

# ============================================================
# STUDENT INFO
# ============================================================
TITLE = "AI-Powered Document Query System"
STUDENT_NAME = "Mr. Dilip Magar"
PRN = "125M1H055"
EMAIL = "dilip.magar25@pccoepune.org"
CONTACT = "9356331879"
GUIDE_NAME = "Dr. Anjana Arakerimath"
COURSE_CODE = "MCA33EL03"
DEPARTMENT = "DEPARTMENT OF MCA"
COLLEGE = "PIMPRI CHINCHWAD COLLEGE OF ENGINEERING SECTOR NO. 26, PRADHIKARAN, NIGDI, PUNE - 44."
YEAR = "2026 - 2027"

# ============================================================
# CONTENT — Condensed for 4-5 pages
# ============================================================

# ------- 1. Introduction -------
INTRODUCTION = (
    "Students and faculty regularly deal with large volumes of departmental PDFs — "
    "syllabi, regulations, lab manuals, and circulars. Finding specific information "
    "from hundreds of pages manually is time-consuming. Existing AI tools like ChatGPT "
    "cannot access private institutional documents and tend to hallucinate when asked "
    "about specific college policies.\n\n"

    "The AI-Powered Document Query System is a full-stack web application that allows "
    "authorized college users to upload departmental PDFs and ask questions in plain English. "
    "It uses Retrieval-Augmented Generation (RAG) to search uploaded documents, find relevant "
    "paragraphs, and generate accurate answers with source citations (file name and page number). "
    "Access is restricted to verified @pccoepune.org Google accounts."
)

# ------- 2. Problem Statement -------
PROBLEM_STATEMENT = (
    "Academic departments manage extensive documentation. When users need specific information, "
    "they must manually search through dozens of PDFs using keyword-based Ctrl+F search, which "
    "fails when different terminology is used. General-purpose LLMs cannot access private documents "
    "and hallucinate incorrect answers. There is no existing tool that combines private document "
    "understanding, hallucination prevention, source citations, and secure access control for "
    "academic use."
)

# ------- 3. Objectives -------
OBJECTIVES = [
    "Build a RAG pipeline that answers questions strictly from uploaded departmental PDFs, preventing hallucination.",
    "Implement multi-stage retrieval using bi-encoder semantic search followed by Cross-Encoder reranking.",
    "Provide verifiable answers by citing the source document name and page number with every response.",
    "Secure the application with Google OAuth authentication restricted to @pccoepune.org users.",
    "Develop a full-stack web application with React frontend and Python FastAPI backend.",
]

# ------- 4. Scope -------
SCOPE = (
    "The proposed system covers the complete pipeline for document-based question "
    "answering within an academic institution. It enables uploading and processing "
    "of multiple PDF documents for text extraction and indexing, converting text "
    "chunks into vector embeddings using Sentence-BERT and storing them in ChromaDB "
    "for efficient retrieval. The system performs semantic search with Cross-Encoder "
    "reranking to accurately identify the most relevant document sections for any "
    "given user query. Answers are generated using the Groq LLM API with strict "
    "prompt guardrails to prevent hallucination, ensuring that the system responds "
    "only from the uploaded document content. Every answer is accompanied by source "
    "citations including the document name and page number for verifiability. Access "
    "to the system is restricted via Google OAuth 2.0 to users with a verified "
    "@pccoepune.org email address, ensuring that only authorized college users can "
    "interact with the application."
)

# ------- 5. Methodology -------
METHODOLOGY_INTRO = (
    "The AI-Powered Document Query System follows a Retrieval-Augmented Generation (RAG) "
    "architecture — it first retrieves relevant sections from uploaded documents and then "
    "uses an AI model to generate answers based only on those sections. The complete system "
    "flow is illustrated in the diagram below."
)

FLOW_DIAGRAM_PATH = os.path.join(
    r"d:\Projects_2026\PDF-RAG-Chatbot",
    "rag_flow_diagram.png"
)

# ------- 6. Requirements -------
FUNCTIONAL_REQUIREMENTS = [
    "Allow upload of PDF files and automatic text extraction and embedding.",
    "Accept natural language questions via a chat interface.",
    "Retrieve relevant sections using semantic search and reranking.",
    "Generate LLM-based answers strictly from uploaded document context.",
    "Display source document name and page number with every answer.",
    "Require Google OAuth authentication; reject non @pccoepune.org emails.",
]

NONFUNCTIONAL_REQUIREMENTS = [
    ("Performance:", "Answer generation within 5–10 seconds."),
    ("Security:", "All data stored locally; no external cloud storage. Google OAuth with domain restriction."),
    ("Usability:", "Clean, intuitive chat interface similar to ChatGPT."),
    ("Scalability:", "New documents can be added without reprocessing existing ones."),
]

# ------- 7. Software & Hardware Requirements -------
SOFTWARE_REQUIREMENTS = [
    ("Frontend", "React 18, Tailwind CSS v3, Vite"),
    ("Backend", "Python 3.10+, FastAPI"),
    ("Embedding Model", "SentenceTransformer (BAAI/bge-small-en-v1.5)"),
    ("Reranker", "CrossEncoder (ms-marco-MiniLM-L-6-v2)"),
    ("Vector DB", "ChromaDB"),
    ("LLM API", "Groq API (Qwen 3-32B)"),
    ("Auth", "Google OAuth 2.0"),
    ("Browser", "Chrome / Firefox (latest)"),
]

TECHNICAL_REQUIREMENTS = [
    ("Hardware:", "8 GB RAM minimum, modern multi-core processor, 10 GB free disk space, stable internet connection."),
    ("OS:", "Windows 10/11, Linux (Ubuntu 20.04+), or macOS."),
]

# ------- 8. Expected Outcomes -------
EXPECTED_OUTCOMES = [
    "A functional web app where authorized users upload PDFs and ask questions in natural language.",
    "Accurate, hallucination-free answers generated strictly from uploaded document content.",
    "Source citations (document name, page number) displayed with every answer.",
    "Secure access restricted to verified @pccoepune.org email addresses.",
]

# ------- 9. Conclusion -------
CONCLUSION = (
    "The AI-Powered Document Query System solves a real problem faced by students and faculty — "
    "quickly finding specific information from departmental PDF collections. By combining "
    "semantic search (Sentence-BERT) with LLM-based answer generation (Groq API) in a RAG "
    "architecture, the system delivers accurate, source-verified answers without hallucination. "
    "Multi-stage retrieval with Cross-Encoder reranking ensures precision, and Google OAuth "
    "restricts access to authorized college users only."
)


# ============================================================
# DOCUMENT GENERATION
# ============================================================

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_blank_lines(doc, count=1):
    for _ in range(count):
        p = doc.add_paragraph()
        run = p.add_run()
        set_font(run, size=6)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def add_heading_text(doc, text, size=13):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=size, bold=True)
    p.space_after = Pt(4)
    p.space_before = Pt(6)
    return p


def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_font(run, size=12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(16)
    return p


def add_bullet_point(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.clear()
    run = p.add_run(text)
    set_font(run, size=12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(16)
    return p


def create_cover_page(doc):
    add_blank_lines(doc, 1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Mini Project Synopsis")
    set_font(run, size=18, bold=True)

    add_blank_lines(doc, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Title of the Project")
    set_font(run, size=12, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'"{TITLE}"')
    set_font(run, size=16, bold=True)

    add_blank_lines(doc, 1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(STUDENT_NAME)
    set_font(run, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"({PRN})")
    set_font(run, size=12)

    add_blank_lines(doc, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(GUIDE_NAME)
    set_font(run, size=12)

    add_blank_lines(doc, 3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DEPARTMENT)
    set_font(run, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COLLEGE)
    set_font(run, size=11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"({YEAR})")
    set_font(run, size=12, bold=True)


def create_index_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Index")
    set_font(run, size=16, bold=True)

    add_blank_lines(doc, 1)

    index_items = [
        "Introduction",
        "Problem Statement",
        "Objectives of Proposed System",
        "Scope of Proposed System",
        "Methodology",
        "Functional Requirements",
        "Non-functional Requirements",
        "Software & Hardware Requirements",
        "Expected Outcomes",
        "Conclusion",
    ]

    for i, item in enumerate(index_items, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"{i}. {item}")
        set_font(run, size=12)
        p.paragraph_format.space_after = Pt(3)


def create_student_details_page(doc):
    add_blank_lines(doc, 2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("STUDENT DETAILS")
    set_font(run, size=14, bold=True)

    add_blank_lines(doc, 2)

    details = [
        ("Name of Student:", STUDENT_NAME.replace("Mr. ", "")),
        ("PRN:", PRN),
        ("Email:", EMAIL),
        ("Contact No.:", CONTACT),
    ]

    for label, value in details:
        p = doc.add_paragraph()
        run_label = p.add_run(f"{label} ")
        set_font(run_label, size=12, bold=True)
        run_value = p.add_run(value)
        set_font(run_value, size=12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(18)


def generate_flow_diagram(output_path):
    """Generate a comprehensive system flow diagram with decision branches."""

    fig, ax = plt.subplots(figsize=(10, 15))
    ax.set_xlim(0, 14)
    ax.set_ylim(3, 23)
    ax.axis('off')
    fig.patch.set_facecolor('white')

    # ---- Constants ----
    BW, BH = 3.0, 1.1
    DW, DH = 3.2, 1.6
    GAP = 0.4
    CX, LX, RX, EX = 6.0, 3.0, 9.0, 12.5

    def draw_box(x, y, text, w=BW, h=BH, fc='white', fs=8.5, bold=False):
        rect = mpatches.FancyBboxPatch(
            (x - w/2, y - h/2), w, h,
            boxstyle="round,pad=0.1",
            facecolor=fc, edgecolor='black', lw=1.1
        )
        ax.add_patch(rect)
        ax.text(x, y, text, ha='center', va='center', fontsize=fs,
                fontfamily='serif', fontweight='bold' if bold else 'normal',
                linespacing=1.3)

    def draw_diamond(x, y, text, w=DW, h=DH, fc='#EDEDED', fs=7.5):
        pts = [(x, y + h/2), (x + w/2, y), (x, y - h/2), (x - w/2, y)]
        ax.add_patch(plt.Polygon(pts, fc=fc, ec='black', lw=1.1, closed=True))
        ax.text(x, y, text, ha='center', va='center', fontsize=fs,
                fontfamily='serif', linespacing=1.2)

    def arr(x1, y1, x2, y2):
        ax.annotate('', xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle='->', color='black', lw=1.1))

    def lbl(x, y, text, fs=7.5):
        ax.text(x, y, text, ha='center', va='center', fontsize=fs,
                fontfamily='serif', fontweight='bold',
                bbox=dict(boxstyle='round,pad=0.06', fc='white', ec='none'))

    # Title
    ax.text(7.0, 22.7,
            "System Flow Diagram \u2014 AI-Powered Document Query System",
            ha='center', va='center', fontsize=10.5,
            fontfamily='serif', fontweight='bold')

    # 1. START
    y = 21.8
    draw_box(CX, y, "User Opens\nWeb Application", fc='#E8E8E8', bold=True)

    # 2. AUTH DIAMOND
    y_auth = y - BH/2 - GAP - DH/2
    arr(CX, y - BH/2, CX, y_auth + DH/2)
    draw_diamond(CX, y_auth, "Authenticated\nvia Google OAuth?\n(@pccoepune.org)")

    # No -> Access Denied
    arr(CX + DW/2, y_auth, EX - 0.8, y_auth)
    draw_box(EX, y_auth, "Access\nDenied", w=1.8, h=0.9, fc='#E0E0E0', fs=8.5, bold=True)
    lbl((CX + DW/2 + EX - 0.8) / 2, y_auth + 0.28, "No")

    # Yes -> Dashboard
    y_dash = y_auth - DH/2 - GAP - BH/2
    arr(CX, y_auth - DH/2, CX, y_dash + BH/2)
    lbl(CX + 0.35, (y_auth - DH/2 + y_dash + BH/2) / 2, "Yes")

    # 3. DASHBOARD
    draw_box(CX, y_dash, "Main Dashboard", fc='#E8E8E8', bold=True)

    # SPLIT
    sb = y_dash - BH/2
    y_col = sb - 0.85
    ax.plot([CX, LX], [sb, sb - 0.25], 'k-', lw=1.1)
    ax.plot([CX, RX], [sb, sb - 0.25], 'k-', lw=1.1)
    arr(LX, sb - 0.25, LX, y_col + BH/2)
    arr(RX, sb - 0.25, RX, y_col + BH/2)
    lbl(LX + 0.9, sb - 0.08, "Upload", fs=7)
    lbl(RX - 0.9, sb - 0.08, "Query", fs=7)
    ax.text(LX, y_col + BH/2 + 0.32, "\u2014 DOCUMENT INGESTION \u2014",
            ha='center', fontsize=7, fontfamily='serif',
            fontweight='bold', color='#666')
    ax.text(RX, y_col + BH/2 + 0.32, "\u2014 QUERY & RESPONSE \u2014",
            ha='center', fontsize=7, fontfamily='serif',
            fontweight='bold', color='#666')

    # LEFT COLUMN
    left_items = [
        "User Uploads\nPDF Documents",
        "Text Extraction\n(PyPDF)",
        "Text Chunking\n(800 chars, 150 overlap)",
        "Embedding Generation\n(Sentence-BERT)",
        "Store in Vector DB\n(ChromaDB)",
    ]
    ly = y_col
    for i, text in enumerate(left_items):
        draw_box(LX, ly, text)
        if i < len(left_items) - 1:
            nly = ly - BH - GAP
            arr(LX, ly - BH/2, LX, nly + BH/2)
            ly = nly

    # RIGHT COLUMN
    right_items = [
        "User Asks\na Question",
        "Query Embedding\n(Sentence-BERT)",
        "Semantic Search\n(ChromaDB \u2014 Top 20)",
        "Cross-Encoder Reranking\n(ms-marco-MiniLM \u2014 Top 5)",
    ]
    ry = y_col
    for i, text in enumerate(right_items):
        draw_box(RX, ry, text)
        if i < len(right_items) - 1:
            nry = ry - BH - GAP
            arr(RX, ry - BH/2, RX, nry + BH/2)
            ry = nry

    # DECISION: Relevant Chunks?
    dy = ry - BH/2 - GAP - DH/2
    arr(RX, ry - BH/2, RX, dy + DH/2)
    draw_diamond(RX, dy, "Relevant\nChunks\nFound?")
    arr(RX + DW/2, dy, EX - 1.0, dy)
    draw_box(EX, dy, "Display\nOut-of-Scope\nMessage", w=2.0, h=1.1, fc='#E0E0E0', fs=7.5, bold=True)
    lbl((RX + DW/2 + EX - 1.0) / 2, dy + 0.28, "No")

    py = dy - DH/2 - GAP - BH/2
    arr(RX, dy - DH/2, RX, py + BH/2)
    lbl(RX + 0.35, (dy - DH/2 + py + BH/2) / 2, "Yes")

    # REMAINING STEPS
    remaining = [
        ("Prompt Construction\n+ Strict Guardrails", BH),
        ("LLM Answer Generation\n(Groq API \u2014 Qwen 3-32B)", BH),
        ("Display Answer +\nSource Citations\n(Document Name, Page No.)", 1.4),
    ]
    for i, (text, h) in enumerate(remaining):
        draw_box(RX, py, text, h=h)
        if i < len(remaining) - 1:
            npy = py - h/2 - GAP - BH/2
            arr(RX, py - h/2, RX, npy + BH/2)
            py = npy

    plt.tight_layout(pad=0.5)
    fig.savefig(output_path, dpi=200, bbox_inches='tight',
                facecolor='white', edgecolor='none')
    plt.close(fig)
    print(f"Flow diagram saved to: {output_path}")


def create_synopsis():

    doc = Document()

    # Default style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ---- COVER PAGE ----
    create_cover_page(doc)
    doc.add_page_break()

    # ---- INDEX PAGE ----
    create_index_page(doc)

    doc.add_page_break()

    # ---- 1. INTRODUCTION ----
    add_heading_text(doc, "1. Introduction")
    for para in INTRODUCTION.split("\n\n"):
        para = para.strip()
        if para:
            add_body_paragraph(doc, para)

    # ---- 2. PROBLEM STATEMENT ----
    add_heading_text(doc, "2. Problem Statement")
    add_body_paragraph(doc, PROBLEM_STATEMENT)

    # ---- 3. OBJECTIVES ----
    add_heading_text(doc, "3. Objectives of Proposed System")
    for obj in OBJECTIVES:
        add_bullet_point(doc, obj)

    # ---- 4. SCOPE ----
    add_heading_text(doc, "4. Scope of Proposed System")
    add_body_paragraph(doc, SCOPE)

    # ---- 5. METHODOLOGY ----
    add_heading_text(doc, "5. Methodology")
    add_body_paragraph(doc, METHODOLOGY_INTRO)

    # Generate and embed flow diagram
    generate_flow_diagram(FLOW_DIAGRAM_PATH)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(FLOW_DIAGRAM_PATH, width=Inches(5.0))
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Fig 1: System Flow Diagram")
    set_font(run, size=10, bold=True, italic=True)
    p.paragraph_format.space_after = Pt(6)

    # ---- 6. FUNCTIONAL REQUIREMENTS ----
    add_heading_text(doc, "6. Functional Requirements")
    for req in FUNCTIONAL_REQUIREMENTS:
        add_bullet_point(doc, req)

    # ---- 7. NON-FUNCTIONAL REQUIREMENTS ----
    add_heading_text(doc, "7. Non-functional Requirements")
    for label, desc in NONFUNCTIONAL_REQUIREMENTS:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_label = p.add_run(label + " ")
        set_font(run_label, size=12, bold=True)
        run_desc = p.add_run(desc)
        set_font(run_desc, size=12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(16)

    # ---- 8. SOFTWARE & HARDWARE REQUIREMENTS ----
    add_heading_text(doc, "8. Software & Hardware Requirements")

    # Software table
    p = doc.add_paragraph()
    run = p.add_run("Software Requirements:")
    set_font(run, size=12, bold=True)
    p.paragraph_format.space_after = Pt(2)

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology"
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, size=11, bold=True)

    for component, tech in SOFTWARE_REQUIREMENTS:
        row_cells = table.add_row().cells
        row_cells[0].text = component
        row_cells[1].text = tech
        for cell in row_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run, size=11)

    # Hardware/OS requirements
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("Hardware & OS Requirements:")
    set_font(run, size=12, bold=True)
    p.paragraph_format.space_after = Pt(2)

    for label, desc in TECHNICAL_REQUIREMENTS:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run_label = p.add_run(label + " ")
        set_font(run_label, size=12, bold=True)
        run_desc = p.add_run(desc)
        set_font(run_desc, size=12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(16)

    # ---- 9. EXPECTED OUTCOMES ----
    add_heading_text(doc, "9. Expected Outcomes")
    for outcome in EXPECTED_OUTCOMES:
        add_bullet_point(doc, outcome)

    # ---- 10. CONCLUSION ----
    add_heading_text(doc, "10. Conclusion")
    add_body_paragraph(doc, CONCLUSION)

    # ---- STUDENT DETAILS PAGE ----
    doc.add_page_break()
    create_student_details_page(doc)

    # ---- SAVE ----
    output_path = os.path.join(
        r"d:\Projects_2026\PDF-RAG-Chatbot",
        "Miniproject_Synopsis_Dilip_Magar.docx"
    )
    doc.save(output_path)
    print(f"Mini Project Synopsis saved to: {output_path}")
    print("Generation complete!")


if __name__ == "__main__":
    create_synopsis()
