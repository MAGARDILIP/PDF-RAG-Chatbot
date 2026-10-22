"""
Seminar Synopsis Generator
Generates the synopsis .docx for MCA Semester IV Seminar
Title: Comparative Study of Transformer Models (BERT vs GPT)
Student: Dilip Magar | PRN: 125M1H055
Guide: Dr. Anjana Arakerimath
"""

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import os

# ============================================================
# CONTENT — All 7 sections written in student-like language
# ============================================================

TITLE = "Comparative Study of Transformer Models (BERT vs GPT)"
STUDENT_NAME = "Mr. Dilip Magar"
PRN = "125M1H055"
GUIDE_NAME = "Dr. Anjana Arakerimath"
DEPARTMENT = "Department of Master of Computer Application"
COLLEGE_LINE1 = "Pimpri Chinchwad Education Trust's"
COLLEGE_LINE2 = "Pimpri Chinchwad College of Engineering"

# ------- Section 1: Abstract (target: 150-200 words) -------
ABSTRACT = (
    "This seminar presents a comparative study of two widely used Transformer-based "
    "language models — BERT (Bidirectional Encoder Representations from Transformers) "
    "and GPT (Generative Pre-trained Transformer). Both of these models are built on "
    "the Transformer architecture proposed by Vaswani et al. in 2017, but they follow "
    "very different design philosophies. BERT uses an encoder-only structure and learns "
    "by predicting masked words in a sentence, which makes it good at understanding "
    "text meaning. GPT, on the other hand, uses a decoder-only structure and learns by "
    "predicting the next word, which makes it suitable for generating text. "
    "In this seminar, the internal architecture, training methods, and practical "
    "applications of both models are studied and compared across Natural Language "
    "Processing (NLP) tasks like text classification, question answering, and text "
    "generation. As a practical demonstration, this study also looks at how both model "
    "families can work together in a Retrieval-Augmented Generation (RAG) system, "
    "where BERT-based models handle document retrieval and reranking, while a GPT-based "
    "model generates the final answer. The expected outcome is to develop a clear "
    "understanding of which model type is better suited for which kind of NLP task."
)

# ------- Section 2: Introduction (target: 150-200 words) -------
INTRODUCTION = (
    "In 2017, Vaswani et al. introduced the Transformer architecture in their paper "
    "\"Attention Is All You Need,\" which changed the field of Natural Language Processing "
    "completely. Before Transformers, models like RNNs and LSTMs were used for language "
    "tasks, but they were slow and struggled with long sentences. The Transformer solved "
    "this by using a mechanism called self-attention, which allows the model to look at "
    "all words in a sentence at once instead of one by one. "
    "After the Transformer was introduced, two major branches emerged. Google released "
    "BERT in 2018, which uses only the encoder part and is designed to understand text "
    "deeply. OpenAI released GPT in the same year, which uses only the decoder part and "
    "is designed to generate text fluently. Today, both models and their successors are "
    "used everywhere — from Google Search to ChatGPT. "
    "I chose this topic because while working on my mini project (InfoFetch — a PDF "
    "chatbot), I used BERT-based models for finding relevant documents and a GPT-based "
    "model for generating answers. This made me curious about what exactly makes these "
    "two models different and why each one is better at certain tasks."
)

# ------- Section 3: Objectives (target: 50-100 words) -------
OBJECTIVES = [
    "To study and understand the internal architecture of BERT (encoder-only) and GPT (decoder-only) Transformer models.",
    "To compare their pre-training strategies — Masked Language Modeling (MLM) in BERT versus autoregressive next-word prediction in GPT.",
    "To evaluate their strengths and weaknesses across different NLP tasks such as text classification, named entity recognition, question answering, and text generation.",
    "To demonstrate their complementary usage through a practical RAG-based case study (InfoFetch project) where BERT handles retrieval and GPT handles generation.",
]

# ------- Section 4: Literature Survey (target: 200-500 words) -------
LITERATURE_SURVEY = (
    "The Transformer architecture was first proposed by Vaswani et al. (2017) in their "
    "landmark paper \"Attention Is All You Need.\" They introduced the self-attention mechanism "
    "which allowed models to process all positions in a sequence simultaneously, removing "
    "the sequential bottleneck of recurrent networks. This architecture became the foundation "
    "for almost all modern NLP models.\n\n"

    "Building on this foundation, Devlin et al. (2019) introduced BERT, which uses only "
    "the encoder portion of the Transformer. BERT is pre-trained using two tasks: Masked "
    "Language Modeling (MLM), where random words in a sentence are hidden and the model "
    "learns to predict them, and Next Sentence Prediction (NSP), where the model learns "
    "whether two sentences follow each other. Because BERT reads text in both directions "
    "(left-to-right and right-to-left simultaneously), it develops a deep understanding of "
    "word context. This made BERT very effective for tasks like sentiment analysis, named "
    "entity recognition, and extractive question answering.\n\n"

    "Around the same time, Radford et al. (2018) at OpenAI introduced GPT, which uses only "
    "the decoder portion of the Transformer. Unlike BERT, GPT reads text in one direction "
    "(left-to-right) and is trained to predict the next word in a sequence. This autoregressive "
    "approach makes GPT naturally good at generating coherent and fluent text. The subsequent "
    "versions — GPT-2 (Radford et al., 2019) and GPT-3 (Brown et al., 2020) — scaled the "
    "model significantly and demonstrated impressive few-shot and zero-shot learning abilities, "
    "meaning the model could perform tasks it was never explicitly trained for.\n\n"

    "Reimers and Gurevych (2019) introduced Sentence-BERT, a modification of BERT that "
    "produces fixed-size sentence embeddings suitable for efficient semantic similarity "
    "comparison. This is particularly useful in information retrieval systems where fast "
    "vector comparison is needed across thousands of documents.\n\n"

    "Lewis et al. (2020) proposed the Retrieval-Augmented Generation (RAG) framework, which "
    "combines retrieval-based methods (typically using BERT-family embeddings) with "
    "generation-based methods (typically using GPT-family models). Their work showed that "
    "this combination significantly reduces hallucination and improves factual accuracy "
    "in knowledge-intensive tasks.\n\n"

    "Most existing studies compare BERT and GPT independently on benchmark datasets. However, "
    "there is limited work that examines how these two model families can complement each "
    "other in a single unified system. This seminar aims to fill that gap by studying both "
    "models individually and then demonstrating their combined usage in a practical RAG "
    "pipeline."
)

# ------- Section 5: Problem Statement (target: 50-100 words) -------
PROBLEM_STATEMENT = (
    "BERT and GPT are both built on the Transformer architecture, yet they serve "
    "fundamentally different purposes — understanding versus generation. Most academic "
    "comparisons treat them as competing models, but in real-world applications like "
    "chatbots and search engines, they often work together. Students and developers "
    "frequently struggle to decide which model to use for a given task. This seminar "
    "addresses the need for a clear, practical comparison that highlights when to use "
    "BERT, when to use GPT, and how they can complement each other."
)

# ------- Section 6: Plan of Study (target: 100-150 words) -------
PLAN_OF_STUDY = [
    ("Phase 1 — Literature Review (Weeks 1–2):",
     "Study the original Transformer paper and the foundational BERT and GPT papers. "
     "Understand the self-attention mechanism and how encoder and decoder architectures differ."),
    ("Phase 2 — Architecture Comparison (Weeks 3–4):",
     "Perform a detailed comparison of BERT's encoder-only and GPT's decoder-only architectures. "
     "Study their pre-training objectives (MLM vs autoregressive) and fine-tuning approaches."),
    ("Phase 3 — Task-wise Performance Analysis (Weeks 5–7):",
     "Compare both models across NLP tasks — classification, question answering, summarization, "
     "and generation — using published benchmarks and research findings."),
    ("Phase 4 — Practical Case Study (Weeks 8–9):",
     "Demonstrate their complementary roles using the InfoFetch RAG project, where Sentence-BERT "
     "handles document retrieval and a GPT-based model handles answer generation."),
    ("Phase 5 — Report Writing and Presentation (Weeks 10–12):",
     "Compile all findings into the final seminar report and prepare for the presentation."),
]

# ------- Section 7: References (IEEE format) -------
REFERENCES = [
    'A. Vaswani et al., "Attention Is All You Need," in Advances in Neural Information Processing Systems (NeurIPS), vol. 30, pp. 5998–6008, 2017.',
    'J. Devlin, M. W. Chang, K. Lee, and K. Toutanova, "BERT: Pre-training of Deep Bidirectional Transformers for Language Understanding," in Proceedings of NAACL-HLT, pp. 4171–4186, 2019.',
    'A. Radford, K. Narasimhan, T. Salimans, and I. Sutskever, "Improving Language Understanding by Generative Pre-Training," OpenAI Technical Report, 2018.',
    'A. Radford et al., "Language Models are Unsupervised Multitask Learners," OpenAI Technical Report, 2019.',
    'T. Brown et al., "Language Models are Few-Shot Learners," in Advances in Neural Information Processing Systems (NeurIPS), vol. 33, pp. 1877–1901, 2020.',
    'P. Lewis et al., "Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks," in Advances in Neural Information Processing Systems (NeurIPS), vol. 33, pp. 9459–9474, 2020.',
    'N. Reimers and I. Gurevych, "Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks," in Proceedings of EMNLP, pp. 3982–3992, 2019.',
    'R. Nogueira and K. Cho, "Passage Re-ranking with BERT," arXiv preprint arXiv:1901.04085, 2019.',
]


# ============================================================
# DOCUMENT GENERATION
# ============================================================

def set_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    """Helper to set font properties on a run."""
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_blank_lines(doc, count=1):
    """Add empty paragraphs for spacing."""
    for _ in range(count):
        p = doc.add_paragraph()
        run = p.add_run()
        set_font(run, size=12)


def create_cover_page(doc):
    """Build the cover page matching the guide's format."""

    # "Synopsis Report" heading
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Synopsis Report")
    set_font(run, size=16, bold=True)

    # "On"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("On")
    set_font(run, size=12)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'"{TITLE}"')
    set_font(run, size=20, bold=True)

    add_blank_lines(doc, 2)

    # Submitted By
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Submitted By")
    set_font(run, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(STUDENT_NAME)
    set_font(run, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"Post Graduate Student with PRN No. {PRN}")
    set_font(run, size=12)

    add_blank_lines(doc, 1)

    # Under the Guidance of
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Under the Guidance of")
    set_font(run, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(GUIDE_NAME)
    set_font(run, size=12)

    add_blank_lines(doc, 3)

    # Department and College
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(DEPARTMENT)
    set_font(run, size=12, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COLLEGE_LINE1)
    set_font(run, size=12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(COLLEGE_LINE2)
    set_font(run, size=12, bold=True)


def add_section_heading(doc, number, title):
    """Add a numbered section heading."""
    p = doc.add_paragraph()
    run = p.add_run(f"{number}. {title}")
    set_font(run, size=14, bold=True)
    p.space_after = Pt(6)


def add_body_text(doc, text):
    """Add a body paragraph with justified alignment."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_font(run, size=12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = Pt(18)  # 1.5 line spacing


def create_synopsis():
    """Main function to generate the complete synopsis document."""
    doc = Document()

    # Set default font for the whole document
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ---- COVER PAGE ----
    create_cover_page(doc)

    # Page break after cover
    doc.add_page_break()

    # ---- CONTENT HEADING ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Seminar Synopsis Content")
    set_font(run, size=16, bold=True)

    add_blank_lines(doc, 1)

    # ---- 1. ABSTRACT ----
    add_section_heading(doc, 1, "Abstract")
    add_body_text(doc, ABSTRACT)
    word_count = len(ABSTRACT.split())
    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {word_count} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 2. INTRODUCTION ----
    add_section_heading(doc, 2, "Introduction")
    add_body_text(doc, INTRODUCTION)
    word_count = len(INTRODUCTION.split())
    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {word_count} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 3. OBJECTIVES ----
    add_section_heading(doc, 3, "Objectives")
    for obj in OBJECTIVES:
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Clear default run and add our own
        p.clear()
        run = p.add_run(obj)
        set_font(run, size=12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(18)

    total_obj_words = sum(len(o.split()) for o in OBJECTIVES)
    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {total_obj_words} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 4. LITERATURE SURVEY ----
    add_section_heading(doc, 4, "Literature Survey")
    # Split by double newlines for separate paragraphs
    for para_text in LITERATURE_SURVEY.split("\n\n"):
        para_text = para_text.strip()
        if para_text:
            add_body_text(doc, para_text)

    word_count = len(LITERATURE_SURVEY.split())
    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {word_count} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 5. PROBLEM STATEMENT ----
    add_section_heading(doc, 5, "Problem Statement")
    add_body_text(doc, PROBLEM_STATEMENT)
    word_count = len(PROBLEM_STATEMENT.split())
    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {word_count} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 6. PLAN OF STUDY ----
    add_section_heading(doc, 6, "Plan of Study")
    p_intro = doc.add_paragraph()
    p_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p_intro.add_run(
        "The seminar will be carried out in the following phases:"
    )
    set_font(run, size=12)
    p_intro.paragraph_format.space_after = Pt(6)

    total_plan_words = len("The seminar will be carried out in the following phases:".split())
    for phase_title, phase_desc in PLAN_OF_STUDY:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        # Bold phase title
        run_title = p.add_run(phase_title + " ")
        set_font(run_title, size=12, bold=True)
        # Normal description
        run_desc = p.add_run(phase_desc)
        set_font(run_desc, size=12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(18)
        total_plan_words += len(phase_title.split()) + len(phase_desc.split())

    p = doc.add_paragraph()
    run = p.add_run(f"(Word count: {total_plan_words} words)")
    set_font(run, size=10, italic=True)

    add_blank_lines(doc, 1)

    # ---- 7. REFERENCES ----
    add_section_heading(doc, 7, "References")
    for i, ref in enumerate(REFERENCES, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(f"[{i}] {ref}")
        set_font(run, size=11)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(18)

    # ---- SAVE ----
    output_path = os.path.join(
        r"d:\Projects_2026\PDF-RAG-Chatbot",
        "Seminar_Synopsis_Dilip_Magar.docx"
    )
    doc.save(output_path)
    print(f"Synopsis saved to: {output_path}")

    # ---- PRINT WORD COUNTS ----
    print("\n--- Word Count Verification ---")
    print(f"Abstract:          {len(ABSTRACT.split())} words  (limit: 150-200)")
    print(f"Introduction:      {len(INTRODUCTION.split())} words  (limit: 150-200)")
    obj_words = sum(len(o.split()) for o in OBJECTIVES)
    print(f"Objectives:        {obj_words} words  (limit: 50-100)")
    print(f"Literature Survey: {len(LITERATURE_SURVEY.split())} words  (limit: 200-500)")
    print(f"Problem Statement: {len(PROBLEM_STATEMENT.split())} words  (limit: 50-100)")
    plan_words = sum(len(t.split()) + len(d.split()) for t, d in PLAN_OF_STUDY) + 10
    print(f"Plan of Study:     {plan_words} words  (limit: 100-150)")
    print(f"References:        {len(REFERENCES)} entries")


if __name__ == "__main__":
    create_synopsis()
